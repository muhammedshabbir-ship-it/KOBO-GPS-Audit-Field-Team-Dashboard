"""
KOBO GPS Audit Dashboard
------------------------
Master file  : master_kobo_file.xlsx  (saved next to this script, can be replaced from the dashboard)
Upload file  : PEP_MA_TRACKING_REPORT ... .xlsb / .xlsx / .csv  (sheet "KOBO RD")

GPS rule  (distance between master LAT/LONG and KOBO GPS)
    < 50 m        -> OK      (green)
    50 - 150 m    -> Review  (pink)
    > 150 m       -> Mismatch(red)
"""
import io
import os
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st

# ----------------------------------------------------------------------------
# CONFIG
# ----------------------------------------------------------------------------
st.set_page_config(page_title="KOBO GPS Audit Dashboard", page_icon="📍", layout="wide")

BASE_DIR = Path(__file__).parent
MASTER_PATH = BASE_DIR / "master_kobo_file.xlsx"
KOBO_SHEET = "KOBO RD"
TZ = "Asia/Karachi"

OK_M, REVIEW_M = 50, 150
S_OK, S_REV, S_RED, S_NA = "OK (<50m)", "Review (50-150m)", "Mismatch (>150m)", "No master / bad GPS"
COLORS = {S_OK: "#22c55e", S_REV: "#f472b6", S_RED: "#ef4444", S_NA: "#94a3b8"}
ORDER = [S_OK, S_REV, S_RED, S_NA]

# columns that are compared with the master (as requested)
KEY_COLS = {
    "Select City Name": "City",
    "Select Auditor Name": "Auditor",
    "Select Store ID": "Store ID",
    "Select Store Name": "Store Name",
    "Channel Type": "Channel",
    "_1.7: GPS Coordinates_latitude": "Lat",
    "_1.7: GPS Coordinates_longitude": "Long",
}
# extra columns used for the additional information
EXTRA_COLS = {
    "start": "start",
    "end": "end",
    "today": "today",
    "username": "Username",
    "deviceid": "Device ID",
    "1.9: Shop Status Code": "Shop Status",
    "_id": "Submission ID",
    "_start-geopoint_latitude": "Start Lat",
    "_start-geopoint_longitude": "Start Long",
    "_1.7: GPS Coordinates_precision": "GPS Accuracy (m)",
    "Q2:  Cooler Availability & Status": "Cooler Availability",
    "2.1.1: How many PEP coolers are in the store?": "PEP Coolers",
    "2.2.1:  How many KO coolers are in the store?": "KO Coolers",
    "2.3.1:  How many Other coolers are in the store?": "Other Coolers",
    "2.1.7: Is the cooler switched ON and operational?": "PEP Cooler ON",
    "MSL": "MSL",
    "MT SHELVES": "MT Shelves",
}
MASTER_COLS = {
    "REPORTED STORE ID": "Store ID",
    "FIN.STORE NAME": "Master Store Name",
    "CITY CODE": "Master City",
    "CHANNEL": "Master Channel",
    "AUDITOR": "Master Auditor",
    "LAT": "Master Lat",
    "LONG": "Master Long",
}

# ----------------------------------------------------------------------------
# STYLE  (dark blue background)
# ----------------------------------------------------------------------------
st.markdown(
    """
<style>
.stApp {background: linear-gradient(160deg,#071533 0%,#0b2559 55%,#0a1c46 100%); color:#e8eefc;}
header[data-testid="stHeader"] {background: transparent;}
.block-container {padding-top: 1.2rem; padding-bottom: 2rem; max-width: 1500px;}
h1,h2,h3,h4,h5,label,p,span,div {color:#e8eefc;}
.title {font-size:30px; font-weight:800; letter-spacing:.3px; margin-bottom:0;}
.subtitle {color:#9db4e6; font-size:14px; margin-bottom:12px;}
.kpi {background: rgba(255,255,255,.07); border:1px solid rgba(255,255,255,.14); border-radius:14px;
      padding:14px 16px; text-align:center; backdrop-filter: blur(4px);}
.kpi .v {font-size:26px; font-weight:800; line-height:1.15;}
.kpi .l {font-size:12px; color:#9db4e6; text-transform:uppercase; letter-spacing:.8px;}
.kpi .s {font-size:12px; color:#c7d6f7;}
.filterbox {background: rgba(255,255,255,.06); border:1px solid rgba(255,255,255,.12);
            border-radius:14px; padding:10px 14px 2px 14px; margin-bottom:14px;}
.note {background: rgba(56,189,248,.12); border-left:4px solid #38bdf8; padding:10px 14px;
       border-radius:6px; font-size:14px;}
.sumbox {background: rgba(255,255,255,.07); border:1px solid rgba(255,255,255,.14);
         border-radius:14px; padding:16px 22px;}
button[data-baseweb="tab"] {color:#c7d6f7;}
</style>
""",
    unsafe_allow_html=True,
)


def show_df(df, **kw):
    try:
        st.dataframe(df, width="stretch", **kw)
    except TypeError:
        st.dataframe(df, use_container_width=True, **kw)


def show_fig(fig, key=None):
    fig.update_layout(
        template="plotly_dark",
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        margin=dict(l=10, r=10, t=45, b=10),
        legend_title_text="",
    )
    try:
        st.plotly_chart(fig, width="stretch", key=key)
    except TypeError:
        st.plotly_chart(fig, use_container_width=True, key=key)


def kpi(col, label, value, sub="", color="#e8eefc"):
    col.markdown(
        f'<div class="kpi"><div class="l">{label}</div>'
        f'<div class="v" style="color:{color}">{value}</div><div class="s">{sub}</div></div>',
        unsafe_allow_html=True,
    )


# ----------------------------------------------------------------------------
# DATA LOADING
# ----------------------------------------------------------------------------
@st.cache_data(show_spinner=False)
def load_master(file_bytes: bytes) -> pd.DataFrame:
    m = pd.read_excel(io.BytesIO(file_bytes))
    m = m.rename(columns=MASTER_COLS)[list(MASTER_COLS.values())]
    m["Store ID"] = m["Store ID"].astype(str).str.strip().str.upper()
    return m.drop_duplicates("Store ID")


def _read_xlsb(path, sheet, wanted):
    from pyxlsb import open_workbook

    with open_workbook(path) as wb:
        with wb.get_sheet(sheet) as sh:
            rows = sh.rows(sparse=True)
            header = {c.c: c.v for c in next(rows) if c.v is not None}
            pick = {}
            for ci, name in header.items():
                if name in wanted and name not in pick.values():
                    pick[ci] = name
            data = []
            for r in rows:
                d = {c.c: c.v for c in r if c.c in pick}
                data.append([d.get(ci) for ci in pick])
    return pd.DataFrame(data, columns=list(pick.values()))


@st.cache_data(show_spinner=False)
def load_kobo(file_bytes: bytes, filename: str) -> pd.DataFrame:
    wanted = set(KEY_COLS) | set(EXTRA_COLS)
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "xlsb":
        with tempfile.NamedTemporaryFile(suffix=".xlsb", delete=False) as tmp:
            tmp.write(file_bytes)
            path = tmp.name
        try:
            df = _read_xlsb(path, KOBO_SHEET, wanted)
        finally:
            os.unlink(path)
    elif ext == "csv":
        df = pd.read_csv(io.BytesIO(file_bytes), usecols=lambda c: c in wanted)
    else:
        df = pd.read_excel(io.BytesIO(file_bytes), sheet_name=KOBO_SHEET, usecols=lambda c: c in wanted)
    return df


def haversine_m(lat1, lon1, lat2, lon2):
    lat1, lon1, lat2, lon2 = map(np.radians, (lat1, lon1, lat2, lon2))
    a = np.sin((lat2 - lat1) / 2) ** 2 + np.cos(lat1) * np.cos(lat2) * np.sin((lon2 - lon1) / 2) ** 2
    return 2 * 6371000 * np.arcsin(np.sqrt(a))


@st.cache_data(show_spinner=False)
def process(kobo: pd.DataFrame, master: pd.DataFrame) -> pd.DataFrame:
    d = kobo.rename(columns={**KEY_COLS, **EXTRA_COLS}).copy()
    d["Store ID"] = d["Store ID"].astype(str).str.strip().str.upper()
    d = d.merge(master, on="Store ID", how="left")

    for c in ["Lat", "Long", "Master Lat", "Master Long", "Start Lat", "Start Long", "GPS Accuracy (m)"]:
        d[c] = pd.to_numeric(d[c], errors="coerce")

    # ---- distance to master (the main check) ----
    good = (
        d["Lat"].between(-90, 90) & d["Long"].between(-180, 180) & (d["Lat"] != 0) & (d["Long"] != 0)
        & d["Master Lat"].notna() & d["Master Long"].notna()
    )
    d["Distance (m)"] = np.where(good, haversine_m(d["Lat"], d["Long"], d["Master Lat"], d["Master Long"]), np.nan)
    d["GPS Status"] = np.select(
        [d["Distance (m)"].isna(), d["Distance (m)"] < OK_M, d["Distance (m)"] <= REVIEW_M],
        [S_NA, S_OK, S_REV],
        default=S_RED,
    )

    # ---- start location vs shop location (was the auditor at the shop when the form started?) ----
    sg = d["Start Lat"].notna() & d["Start Long"].notna() & good
    d["Start Dist (m)"] = np.where(sg, haversine_m(d["Start Lat"], d["Start Long"], d["Master Lat"], d["Master Long"]), np.nan)

    # ---- master vs kobo attribute checks ----
    d["Auditor Mismatch"] = d["Master Auditor"].notna() & (d["Auditor"] != d["Master Auditor"])
    d["City Mismatch"] = d["Master City"].notna() & (d["City"] != d["Master City"])
    d["Channel Mismatch"] = d["Master Channel"].notna() & (d["Channel"] != d["Master Channel"])

    # ---- time ----
    s = pd.to_datetime(d["start"], format="ISO8601", utc=True, errors="coerce").dt.tz_convert(TZ)
    e = pd.to_datetime(d["end"], format="ISO8601", utc=True, errors="coerce").dt.tz_convert(TZ)
    d["Start Time"], d["End Time"] = s, e
    dur = (e - s).dt.total_seconds() / 60
    d["Duration (min)"] = dur.where((dur >= 0) & (dur <= 180))
    d["Date"] = pd.to_datetime(d["today"], errors="coerce").dt.normalize()
    d["Date"] = d["Date"].fillna(s.dt.tz_localize(None).dt.normalize())
    d["Hour"] = s.dt.hour
    d["Fast Interview (<2 min)"] = d["Duration (min)"] < 2

    # ---- duplicate GPS (same coordinates submitted for different stores) ----
    key = d["Lat"].round(5).astype(str) + "," + d["Long"].round(5).astype(str)
    d["Duplicate GPS"] = key.map(d.groupby(key)["Store ID"].nunique()).gt(1) & good

    # ---- device label ----
    d["Device ID"] = d["Device ID"].astype(str)
    d["Device"] = d["Device ID"].str.replace("collect:", "", regex=False).str[:14]

    # ---- cooler / MSL helpers ----
    for c in ["PEP Coolers", "KO Coolers", "Other Coolers", "MSL", "MT Shelves"]:
        d[c] = pd.to_numeric(d[c], errors="coerce")
    return d


# ----------------------------------------------------------------------------
# EXCEL EXPORT (green / pink / red rows)
# ----------------------------------------------------------------------------
def to_colored_excel(df: pd.DataFrame) -> bytes:
    from openpyxl.styles import PatternFill

    out = df.copy()
    for c in out.columns:
        if str(out[c].dtype).startswith("datetime64") and getattr(out[c].dt, "tz", None) is not None:
            out[c] = out[c].dt.tz_localize(None)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        out.to_excel(xw, index=False, sheet_name="GPS_Compare")
        ws = xw.sheets["GPS_Compare"]
        fills = {
            S_OK: PatternFill("solid", fgColor="C6EFCE"),
            S_REV: PatternFill("solid", fgColor="FFC0E0"),
            S_RED: PatternFill("solid", fgColor="FF9999"),
        }
        st_idx = list(out.columns).index("GPS Status") + 1
        for r in range(2, ws.max_row + 1):
            f = fills.get(ws.cell(r, st_idx).value)
            if f:
                for c in range(1, ws.max_column + 1):
                    ws.cell(r, c).fill = f
        ws.freeze_panes = "A2"
    return buf.getvalue()


def style_status(df, col="GPS Status"):
    cmap = {S_OK: "rgba(34,197,94,.30)", S_REV: "rgba(244,114,182,.38)", S_RED: "rgba(239,68,68,.45)"}
    pd.set_option("styler.render.max_elements", 5_000_000)
    return df.style.apply(lambda r: [f"background-color:{cmap.get(r[col], 'transparent')}"] * len(r), axis=1)


# ----------------------------------------------------------------------------
# HEADER + MASTER / UPLOAD
# ----------------------------------------------------------------------------
st.markdown('<div class="title">📍 KOBO GPS Audit & Field Team Dashboard</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="subtitle">Master vs KOBO GPS comparison · Device audit · Team productivity · Cooler & MSL insights</div>',
    unsafe_allow_html=True,
)

with st.expander("📂 Files (master is saved in this app – upload the second-half KOBO file here)", expanded=True):
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("**1) Master file**")
        up_master = st.file_uploader("Replace master (optional)", type=["xlsx"], key="mst")
        if up_master is not None:
            MASTER_PATH.write_bytes(up_master.getvalue())
            load_master.clear()
            st.success("Master saved. It will be used every time the app starts.")
        if MASTER_PATH.exists():
            st.caption(f"Using saved master: `{MASTER_PATH.name}`")
    with c2:
        st.markdown("**2) KOBO tracking file** (sheet: `KOBO RD`)")
        up_kobo = st.file_uploader("Upload .xlsb / .xlsx / .csv", type=["xlsb", "xlsx", "csv"], key="kb")

if not MASTER_PATH.exists():
    st.warning("Master file is not saved yet. Please upload it above.")
    st.stop()

default_path = os.environ.get("KOBO_DEFAULT_PATH")  # optional: auto-load a file for testing
if up_kobo is None and not default_path:
    st.info("⬆️ Upload the KOBO tracking file to see the dashboard.")
    st.stop()

with st.spinner("Reading files… (a large .xlsb can take a minute the first time)"):
    master = load_master(MASTER_PATH.read_bytes())
    if up_kobo is not None:
        kobo = load_kobo(up_kobo.getvalue(), up_kobo.name)
    else:
        kobo = load_kobo(Path(default_path).read_bytes(), default_path)
    data = process(kobo, master)

# ----------------------------------------------------------------------------
# FILTERS (top)
# ----------------------------------------------------------------------------
st.markdown('<div class="filterbox">', unsafe_allow_html=True)
f1, f2, f3, f4, f5, f6 = st.columns([1.1, 1.3, 1, 1.3, 1.2, 1.3])
city = f1.multiselect("City", sorted(data["City"].dropna().unique()))
aud = f2.multiselect("Auditor", sorted(data["Auditor"].dropna().unique()))
chan = f3.multiselect("Channel", sorted(data["Channel"].dropna().unique()))
dmin, dmax = data["Date"].min().date(), data["Date"].max().date()
drange = f4.date_input("Date range", (dmin, dmax), min_value=dmin, max_value=dmax)
gstat = f5.multiselect("GPS status", ORDER)
dev = f6.multiselect("Device", sorted(data["Device"].unique()))
st.markdown("</div>", unsafe_allow_html=True)

f = data.copy()
if city: f = f[f["City"].isin(city)]
if aud: f = f[f["Auditor"].isin(aud)]
if chan: f = f[f["Channel"].isin(chan)]
if gstat: f = f[f["GPS Status"].isin(gstat)]
if dev: f = f[f["Device"].isin(dev)]
if isinstance(drange, (list, tuple)) and len(drange) == 2:
    f = f[(f["Date"].dt.date >= drange[0]) & (f["Date"].dt.date <= drange[1])]
if f.empty:
    st.warning("No data for the selected filters.")
    st.stop()

# ----------------------------------------------------------------------------
# KPIs
# ----------------------------------------------------------------------------
n = len(f)
cnt = f["GPS Status"].value_counts()
pct = lambda k: cnt.get(k, 0) / n * 100
k = st.columns(8)
kpi(k[0], "Total Visits", f"{n:,}", f"{f['Store ID'].nunique():,} unique stores")
kpi(k[1], "Auditors", f["Auditor"].nunique(), f"{f['Device ID'].nunique()} devices")
kpi(k[2], "GPS OK", f"{cnt.get(S_OK, 0):,}", f"{pct(S_OK):.1f}% · < 50 m", COLORS[S_OK])
kpi(k[3], "GPS Review", f"{cnt.get(S_REV, 0):,}", f"{pct(S_REV):.1f}% · 50–150 m", COLORS[S_REV])
kpi(k[4], "GPS Mismatch", f"{cnt.get(S_RED, 0):,}", f"{pct(S_RED):.1f}% · > 150 m", COLORS[S_RED])
kpi(k[5], "Median Distance", f"{f['Distance (m)'].median():.1f} m", f"mean {f['Distance (m)'].mean():.0f} m")
kpi(k[6], "Avg Interview", f"{f['Duration (min)'].mean():.1f} min", f"{int(f['Fast Interview (<2 min)'].sum())} under 2 min")
kpi(k[7], "Completed", f"{(f['Shop Status'] == 'Completed Interview').mean() * 100:.1f}%", f"{int((f['Shop Status'] != 'Completed Interview').sum())} closed")
st.write("")

tabs = st.tabs([
    "📊 Overview", "📍 GPS Compare", "👥 Team Performance", "📱 Device Audit",
    "⏱️ Time & Productivity", "🧊 Coolers & MSL", "🗺️ Map", "📝 Summary", "📥 Data & Download",
])

# ----------------------------------------------------------------------------
# helper tables
# ----------------------------------------------------------------------------
def status_table(by):
    t = f.groupby([by, "GPS Status"]).size().unstack(fill_value=0).reindex(columns=ORDER, fill_value=0)
    t["Total"] = t.sum(axis=1)
    t["% OK"] = (t[S_OK] / t["Total"] * 100).round(1)
    t["% Review"] = (t[S_REV] / t["Total"] * 100).round(1)
    t["% Mismatch"] = (t[S_RED] / t["Total"] * 100).round(1)
    return t.reset_index()


aud_day = f.groupby(["Auditor", "Date"]).agg(
    Visits=("Store ID", "size"), First=("Start Time", "min"), Last=("End Time", "max")
).reset_index()
aud_day["Hours"] = ((aud_day["Last"] - aud_day["First"]).dt.total_seconds() / 3600).clip(upper=16)

# ----------------------------------------------------------------------------
# TAB 1 - OVERVIEW
# ----------------------------------------------------------------------------
with tabs[0]:
    c1, c2 = st.columns([1, 2])
    with c1:
        pie_df = cnt.reindex(ORDER).dropna().rename_axis("GPS Status").reset_index(name="Visits")
        show_fig(px.pie(pie_df, names="GPS Status", values="Visits", hole=0.55, color="GPS Status",
                        color_discrete_map=COLORS, title="GPS check result"), "ov_pie")
    with c2:
        daily = f.groupby(["Date", "GPS Status"]).size().reset_index(name="Visits")
        show_fig(px.bar(daily, x="Date", y="Visits", color="GPS Status", color_discrete_map=COLORS,
                        category_orders={"GPS Status": ORDER}, title="Daily visits by GPS result"), "ov_daily")

    c3, c4 = st.columns(2)
    with c3:
        t = status_table("City")
        show_fig(px.bar(t.melt(id_vars="City", value_vars=ORDER, var_name="GPS Status", value_name="Visits"),
                        x="City", y="Visits", color="GPS Status", color_discrete_map=COLORS,
                        title="City-wise visits & GPS result"), "ov_city")
    with c4:
        t = status_table("Channel")
        show_fig(px.bar(t.melt(id_vars="Channel", value_vars=ORDER, var_name="GPS Status", value_name="Visits"),
                        x="Channel", y="Visits", color="GPS Status", color_discrete_map=COLORS,
                        title="Channel-wise visits & GPS result"), "ov_chan")

# ----------------------------------------------------------------------------
# TAB 2 - GPS COMPARE
# ----------------------------------------------------------------------------
with tabs[1]:
    st.markdown(
        '<div class="note">Distance = straight-line (haversine) distance between the <b>master LAT/LONG</b> and the '
        '<b>KOBO “1.7 GPS Coordinates”</b> of the same Store ID. '
        '🟢 &lt; 50 m OK · 🩷 50–150 m Review · 🔴 &gt; 150 m Mismatch.</div>',
        unsafe_allow_html=True,
    )
    st.write("")
    c1, c2 = st.columns(2)
    with c1:
        h = f[f["Distance (m)"].notna()]
        show_fig(px.histogram(h[h["Distance (m)"] <= 300], x="Distance (m)", nbins=60, color="GPS Status",
                              color_discrete_map=COLORS, title="Distance distribution (0-300 m shown)"), "gps_hist")
    with c2:
        tc = status_table("City").sort_values("% Mismatch", ascending=False)
        show_fig(px.bar(tc, x="City", y=["% Review", "% Mismatch"], barmode="stack",
                        color_discrete_map={"% Review": COLORS[S_REV], "% Mismatch": COLORS[S_RED]},
                        title="% Review + % Mismatch by city"), "gps_city")

    st.markdown("#### Auditor-wise GPS result")
    ta = status_table("Auditor").sort_values("% Mismatch", ascending=False)
    show_df(ta, hide_index=True)

    st.markdown("#### 🔴 Mismatch & 🩷 Review stores (largest distance first)")
    cols = ["GPS Status", "Distance (m)", "City", "Auditor", "Store ID", "Store Name", "Channel",
            "Lat", "Long", "Master Lat", "Master Long", "Date", "Device"]
    bad = f[f["GPS Status"].isin([S_RED, S_REV])].sort_values("Distance (m)", ascending=False)[cols]
    bad = bad.assign(**{"Distance (m)": bad["Distance (m)"].round(1), "Date": bad["Date"].dt.date})
    show_df(style_status(bad), hide_index=True, height=420)

    st.markdown("#### Other master-vs-KOBO checks")
    m1, m2, m3, m4 = st.columns(4)
    kpi(m1, "Auditor ≠ master", int(f["Auditor Mismatch"].sum()))
    kpi(m2, "City ≠ master", int(f["City Mismatch"].sum()))
    kpi(m3, "Channel ≠ master", int(f["Channel Mismatch"].sum()))
    kpi(m4, "Same GPS, different stores", int(f["Duplicate GPS"].sum()), "possible copy-paste GPS")
    mm = f[f["Auditor Mismatch"] | f["City Mismatch"] | f["Channel Mismatch"]][
        ["Store ID", "Store Name", "City", "Master City", "Channel", "Master Channel", "Auditor", "Master Auditor"]]
    if len(mm):
        with st.expander(f"Show {len(mm)} attribute mismatches"):
            show_df(mm, hide_index=True)
    dup = f[f["Duplicate GPS"]].sort_values(["Lat", "Long"])[["Lat", "Long", "Store ID", "Store Name", "Auditor", "Date"]]
    if len(dup):
        with st.expander(f"Show {len(dup)} visits sharing identical GPS with another store"):
            show_df(dup.assign(Date=dup["Date"].dt.date), hide_index=True)

# ----------------------------------------------------------------------------
# TAB 3 - TEAM PERFORMANCE
# ----------------------------------------------------------------------------
with tabs[2]:
    perf = f.groupby("Auditor").agg(
        City=("City", lambda x: ", ".join(sorted(x.dropna().unique()))),
        Visits=("Store ID", "size"),
        Stores=("Store ID", "nunique"),
        Active_Days=("Date", "nunique"),
        Avg_Duration_min=("Duration (min)", "mean"),
        Median_Dist_m=("Distance (m)", "median"),
        Fast_Interviews=("Fast Interview (<2 min)", "sum"),
        Closed_Shops=("Shop Status", lambda x: (x != "Completed Interview").sum()),
        Devices=("Device ID", "nunique"),
    ).reset_index()
    perf["Visits/Day"] = (perf["Visits"] / perf["Active_Days"]).round(1)
    st_ = status_table("Auditor")[["Auditor", "% OK", "% Review", "% Mismatch"]]
    perf = perf.merge(st_, on="Auditor").sort_values("Visits", ascending=False)
    perf[["Avg_Duration_min", "Median_Dist_m"]] = perf[["Avg_Duration_min", "Median_Dist_m"]].round(1)

    c1, c2 = st.columns(2)
    with c1:
        show_fig(px.bar(perf, x="Visits", y="Auditor", orientation="h", color="City",
                        title="Visits per auditor").update_yaxes(categoryorder="total ascending"), "tp_vis")
    with c2:
        show_fig(px.bar(perf.sort_values("Visits/Day"), x="Visits/Day", y="Auditor", orientation="h",
                        color="% Mismatch", color_continuous_scale="Reds",
                        title="Average visits per active day (colour = % GPS mismatch)"), "tp_vpd")
    show_df(perf, hide_index=True)

    piv = f.pivot_table(index="Auditor", columns=f["Date"].dt.strftime("%d-%b"), values="Store ID",
                        aggfunc="count", fill_value=0)
    piv = piv[sorted(piv.columns, key=lambda x: pd.to_datetime(x + "-2026", format="%d-%b-%Y"))]
    show_fig(px.imshow(piv, aspect="auto", color_continuous_scale="Blues", text_auto=True,
                       title="Auditor x Date - number of visits"), "tp_heat")

# ----------------------------------------------------------------------------
# TAB 4 - DEVICE AUDIT
# ----------------------------------------------------------------------------
with tabs[3]:
    st.markdown('<div class="note">Rule checked: <b>one auditor should use only one device</b>, and one device '
                'should belong to only one auditor.</div>', unsafe_allow_html=True)
    st.write("")
    ad = f.groupby("Auditor").agg(
        Devices=("Device ID", "nunique"),
        Device_List=("Device", lambda x: ", ".join(sorted(x.unique()))),
        Visits=("Store ID", "size"),
    ).reset_index()
    ad["Result"] = np.where(ad["Devices"] > 1, "⚠️ Multiple devices", "✅ One device")
    da = f.groupby("Device").agg(
        Auditors=("Auditor", "nunique"),
        Auditor_List=("Auditor", lambda x: ", ".join(sorted(x.unique()))),
        Visits=("Store ID", "size"),
        First_Seen=("Date", "min"),
        Last_Seen=("Date", "max"),
    ).reset_index()
    da["Result"] = np.where(da["Auditors"] > 1, "⚠️ Shared device", "✅ One auditor")
    da["First_Seen"], da["Last_Seen"] = da["First_Seen"].dt.date, da["Last_Seen"].dt.date

    d1, d2, d3, d4 = st.columns(4)
    kpi(d1, "Auditors", len(ad))
    kpi(d2, "Devices", len(da))
    kpi(d3, "Auditors with >1 device", int((ad["Devices"] > 1).sum()), color="#f472b6")
    kpi(d4, "Devices used by >1 auditor", int((da["Auditors"] > 1).sum()), color="#ef4444")
    st.write("")

    sty = lambda df_: df_.style.apply(
        lambda r: ["background-color:rgba(244,114,182,.35)" if "⚠️" in str(r["Result"]) else ""] * len(r), axis=1)
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("#### Auditor → devices")
        show_df(sty(ad.sort_values("Devices", ascending=False)), hide_index=True)
    with c2:
        st.markdown("#### Device → auditors")
        show_df(sty(da.sort_values("Auditors", ascending=False)), hide_index=True)

    multi = ad[ad["Devices"] > 1]["Auditor"]
    if len(multi):
        sub = f[f["Auditor"].isin(multi)].groupby(["Auditor", "Device"]).agg(
            Visits=("Store ID", "size"), First=("Date", "min"), Last=("Date", "max")).reset_index()
        sub["First"], sub["Last"] = sub["First"].dt.date, sub["Last"].dt.date
        show_fig(px.bar(sub, x="Auditor", y="Visits", color="Device", title="Visits by device for auditors using multiple devices"), "dv_multi")
        with st.expander("Date range of each device for those auditors"):
            show_df(sub, hide_index=True)

# ----------------------------------------------------------------------------
# TAB 5 - TIME & PRODUCTIVITY
# ----------------------------------------------------------------------------
with tabs[4]:
    c1, c2 = st.columns(2)
    with c1:
        hr = f.groupby("Hour").size().reset_index(name="Visits")
        show_fig(px.bar(hr, x="Hour", y="Visits", title="Visits by hour of day (start time)"), "tm_hour")
    with c2:
        dd = f["Duration (min)"].dropna()
        show_fig(px.histogram(dd[dd <= 40], nbins=40, title="Interview duration (min, up to 40 shown)",
                              labels={"value": "minutes"}), "tm_dur")

    day = f.groupby("Date").agg(Visits=("Store ID", "size"), Auditors=("Auditor", "nunique")).reset_index()
    day["Visits per Auditor"] = (day["Visits"] / day["Auditors"]).round(1)
    show_fig(px.line(day, x="Date", y=["Visits", "Auditors"], markers=True, title="Daily visits & active auditors"), "tm_day")

    st.markdown("#### Auditor × day: first start, last end and working span")
    ad_show = aud_day.copy()
    ad_show["Date"] = ad_show["Date"].dt.date
    ad_show["First"] = ad_show["First"].dt.strftime("%H:%M")
    ad_show["Last"] = ad_show["Last"].dt.strftime("%H:%M")
    ad_show["Hours"] = ad_show["Hours"].round(1)
    ad_show["Visits/Hour"] = (ad_show["Visits"] / ad_show["Hours"].replace(0, np.nan)).round(1)
    show_df(ad_show.sort_values(["Date", "Auditor"]), hide_index=True, height=380)

    fast = f[f["Fast Interview (<2 min)"]]
    st.markdown(f"#### ⚡ Very fast interviews (< 2 min): {len(fast)}")
    if len(fast):
        fa = fast.groupby("Auditor").size().reset_index(name="Fast interviews").sort_values("Fast interviews", ascending=False)
        show_df(fa, hide_index=True)

# ----------------------------------------------------------------------------
# TAB 6 - COOLERS & MSL
# ----------------------------------------------------------------------------
with tabs[5]:
    done = f[f["Shop Status"] == "Completed Interview"]
    c1, c2, c3, c4 = st.columns(4)
    kpi(c1, "Stores with PEP cooler", f"{(done['PEP Coolers'] > 0).mean() * 100:.1f}%", f"{int((done['PEP Coolers'] > 0).sum())} stores")
    kpi(c2, "Stores with KO cooler", f"{(done['KO Coolers'] > 0).mean() * 100:.1f}%", f"{int((done['KO Coolers'] > 0).sum())} stores")
    kpi(c3, "PEP cooler switched ON", f"{(done['PEP Cooler ON'] == 'Yes').sum() / max(done['PEP Cooler ON'].notna().sum(), 1) * 100:.1f}%",
        f"{int((done['PEP Cooler ON'] == 'No').sum())} switched off")
    kpi(c4, "MSL compliance flag = 1", f"{done['MSL'].mean() * 100:.1f}%", f"{int((done['MSL'] == 0).sum())} stores = 0")
    st.write("")

    a, b = st.columns(2)
    with a:
        sc = f["Shop Status"].value_counts().reset_index()
        sc.columns = ["Shop Status", "Visits"]
        show_fig(px.pie(sc, names="Shop Status", values="Visits", hole=.5, title="Shop status"), "cl_status")
    with b:
        ca = done["Cooler Availability"].fillna("No cooler info").value_counts().reset_index()
        ca.columns = ["Cooler Availability", "Stores"]
        show_fig(px.bar(ca, x="Stores", y="Cooler Availability", orientation="h", title="Cooler availability (who has coolers)"), "cl_avail")

    cc = done.groupby("City").agg(
        PEP=("PEP Coolers", "sum"), KO=("KO Coolers", "sum"), Other=("Other Coolers", "sum")).reset_index()
    show_fig(px.bar(cc, x="City", y=["PEP", "KO", "Other"], barmode="group", title="Total coolers found by city"), "cl_city")

    ms = done.groupby("City").agg(Stores=("Store ID", "size"), MSL_pct=("MSL", lambda x: round(x.mean() * 100, 1)),
                                  PEP_off=("PEP Cooler ON", lambda x: int((x == "No").sum()))).reset_index()
    ms2 = done.groupby("Auditor").agg(Stores=("Store ID", "size"), MSL_pct=("MSL", lambda x: round(x.mean() * 100, 1))).reset_index()
    x1, x2 = st.columns(2)
    with x1:
        st.markdown("#### City: MSL % and PEP coolers switched OFF")
        show_df(ms, hide_index=True)
    with x2:
        st.markdown("#### Auditor: MSL %")
        show_df(ms2.sort_values("MSL_pct"), hide_index=True)

# ----------------------------------------------------------------------------
# TAB 7 - MAP
# ----------------------------------------------------------------------------
with tabs[6]:
    mp = f[f["Lat"].between(20, 40) & f["Long"].between(60, 80)].copy()
    mp["Distance (m)"] = mp["Distance (m)"].round(1)
    kw = dict(lat="Lat", lon="Long", color="GPS Status", color_discrete_map=COLORS,
              category_orders={"GPS Status": ORDER}, hover_name="Store Name",
              hover_data={"Store ID": True, "Auditor": True, "Distance (m)": True, "Lat": False, "Long": False},
              zoom=4.6, height=650)
    if hasattr(px, "scatter_map"):
        fig = px.scatter_map(mp, map_style="carto-darkmatter", **kw)
    else:
        fig = px.scatter_mapbox(mp, mapbox_style="carto-darkmatter", **kw)
    show_fig(fig, "map")

    st.markdown("#### Zoom to one city")
    only = st.checkbox("Show only 🩷 / 🔴 points", value=False)
    mm_ = mp[mp["GPS Status"].isin([S_REV, S_RED])] if only else mp
    cs = st.selectbox("City", sorted(mm_["City"].unique()))
    sub = mm_[mm_["City"] == cs]
    if hasattr(px, "scatter_map"):
        fig2 = px.scatter_map(sub, map_style="carto-darkmatter", **{**kw, "zoom": 10, "height": 550})
    else:
        fig2 = px.scatter_mapbox(sub, mapbox_style="carto-darkmatter", **{**kw, "zoom": 10, "height": 550})
    show_fig(fig2, "map2")

# ----------------------------------------------------------------------------
# TAB 8 - SUMMARY (auto generated)
# ----------------------------------------------------------------------------
with tabs[7]:
    tc = status_table("City")
    tc_ok = tc[tc["Total"] >= 20]
    worst_city = tc_ok.sort_values("% Mismatch", ascending=False).iloc[0]
    best_city = tc_ok.sort_values("% OK", ascending=False).iloc[0]
    ta = status_table("Auditor")
    ta_ok = ta[ta["Total"] >= 20]
    worst_aud = ta_ok.sort_values("% Mismatch", ascending=False).iloc[0]
    best_aud = ta_ok.sort_values("% OK", ascending=False).iloc[0]
    top_aud = f["Auditor"].value_counts()
    dayv = f.groupby("Date").size()
    multi_aud = ad[ad["Devices"] > 1]
    shared_dev = da[da["Auditors"] > 1]
    done = f[f["Shop Status"] == "Completed Interview"]

    lines = [
        f"### Fieldwork overview",
        f"- **{n:,} visits** covering **{f['Store ID'].nunique():,} stores** in **{f['City'].nunique()} cities** by "
        f"**{f['Auditor'].nunique()} auditors** between **{f['Date'].min():%d %b}** and **{f['Date'].max():%d %b %Y}** "
        f"({f['Date'].nunique()} working days, about {n / f['Date'].nunique():.0f} visits/day).",
        f"- Busiest day: **{dayv.idxmax():%d %b}** ({dayv.max()} visits); quietest day: **{dayv.idxmin():%d %b}** ({dayv.min()} visits).",
        f"- Most visits: **{top_aud.index[0]}** ({top_aud.iloc[0]}); fewest: **{top_aud.index[-1]}** ({top_aud.iloc[-1]}).",
        f"- Shop status: {(f['Shop Status'] == 'Completed Interview').mean() * 100:.1f}% completed interviews; "
        f"{int((f['Shop Status'] == 'Temporary Closed').sum())} temporarily closed and {int((f['Shop Status'] == 'Permanently Closed').sum())} permanently closed.",
        "",
        "### GPS check against master",
        f"- 🟢 **{cnt.get(S_OK, 0):,} ({pct(S_OK):.1f}%)** are within 50 m, 🩷 **{cnt.get(S_REV, 0):,} ({pct(S_REV):.1f}%)** are 50-150 m "
        f"and 🔴 **{cnt.get(S_RED, 0):,} ({pct(S_RED):.1f}%)** are above 150 m. Median distance is **{f['Distance (m)'].median():.1f} m**.",
        f"- Highest mismatch city: **{worst_city['City']}** ({worst_city['% Mismatch']}% red, {worst_city['% Review']}% pink). "
        f"Best city: **{best_city['City']}** ({best_city['% OK']}% OK).",
        f"- Highest mismatch auditor: **{worst_aud['Auditor']}** ({worst_aud['% Mismatch']}% red). "
        f"Best auditor: **{best_aud['Auditor']}** ({best_aud['% OK']}% OK).",
        f"- {int(f['Duplicate GPS'].sum())} visits share exactly the same GPS point with another store; "
        f"{int(f['Auditor Mismatch'].sum())} visits have a different auditor than master, {int(f['City Mismatch'].sum())} a different city "
        f"and {int(f['Channel Mismatch'].sum())} a different channel.",
        "",
        "### Device audit",
        f"- **{len(da)} devices** were used by **{len(ad)} auditors**. "
        + (f"**{len(multi_aud)} auditor(s) used more than one device:** " + "; ".join(
            f"{r.Auditor} ({r.Devices})" for r in multi_aud.itertuples()) + "."
           if len(multi_aud) else "Every auditor used only one device."),
        f"- " + (f"**{len(shared_dev)} device(s) were shared by more than one auditor.**" if len(shared_dev)
                 else "No device was shared between two auditors."),
        "",
        "### Time & productivity",
        f"- Average interview length **{f['Duration (min)'].mean():.1f} min** (median {f['Duration (min)'].median():.1f}). "
        f"**{int(f['Fast Interview (<2 min)'].sum())}** interviews were shorter than 2 minutes.",
        f"- Peak working hour: **{f.groupby('Hour').size().idxmax()}:00**. Average auditor works "
        f"**{aud_day['Hours'].mean():.1f} h/day** with **{aud_day['Visits'].mean():.1f} visits/day**.",
        "",
        "### Coolers & MSL (completed interviews)",
        f"- PEP cooler present in **{(done['PEP Coolers'] > 0).mean() * 100:.1f}%** of stores, KO in **{(done['KO Coolers'] > 0).mean() * 100:.1f}%**; "
        f"**{int((done['PEP Cooler ON'] == 'No').sum())}** PEP coolers were found switched off.",
        f"- MSL compliance flag is positive in **{done['MSL'].mean() * 100:.1f}%** of stores.",
    ]
    st.markdown('<div class="sumbox">', unsafe_allow_html=True)
    st.markdown("\n".join(lines))
    st.markdown("</div>", unsafe_allow_html=True)

# ----------------------------------------------------------------------------
# TAB 9 - DATA & DOWNLOAD
# ----------------------------------------------------------------------------
with tabs[8]:
    out_cols = ["GPS Status", "Distance (m)", "City", "Auditor", "Store ID", "Store Name", "Channel", "Lat", "Long",
                "Master Lat", "Master Long", "Master Store Name", "Master City", "Master Channel", "Master Auditor",
                "Auditor Mismatch", "City Mismatch", "Channel Mismatch", "Duplicate GPS", "Start Dist (m)",
                "GPS Accuracy (m)", "Device ID", "Date", "Start Time", "End Time", "Duration (min)", "Shop Status",
                "PEP Coolers", "KO Coolers", "Other Coolers", "PEP Cooler ON", "MSL", "MT Shelves"]
    res = f[out_cols].copy()
    res["Distance (m)"] = res["Distance (m)"].round(1)
    res["Start Dist (m)"] = res["Start Dist (m)"].round(1)
    st.download_button(
        "📥 Download coloured Excel (green / pink / red)",
        data=to_colored_excel(res),
        file_name="GPS_compare_result.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    show_df(style_status(res.assign(Date=res["Date"].dt.date, **{"Start Time": res["Start Time"].dt.strftime("%d-%b %H:%M"),
                                                                "End Time": res["End Time"].dt.strftime("%d-%b %H:%M")})),
            hide_index=True, height=560)
